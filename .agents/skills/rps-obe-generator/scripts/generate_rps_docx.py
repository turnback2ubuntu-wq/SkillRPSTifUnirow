#!/usr/bin/env python3
"""
generate_rps_docx.py
--------------------
Generator dokumen Rencana Pembelajaran Semester (RPS) berbasis Outcome-Based Education (OBE)
yang mengimplementasikan alur Buku Saku RPS OBE dan menghasilkan file .docx berstandar
Format RPS UNIROW (Universitas PGRI Ronggolawe).

Penggunaan:
    python3 generate_rps_docx.py --input data_rps.json --output RPS_MK08_RPL.docx
    python3 generate_rps_docx.py --example
"""

import sys
import os
import json
import argparse
from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_font(cell, name="Calibri", size_pt=9, bold=False, color_rgb=(0, 0, 0), align=WD_ALIGN_PARAGRAPH.LEFT):
    """Mengatur font, ukuran, ketebalan, dan perataan teks pada sebuah cell tabel."""
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            r.font.name = name
            r.font.size = Pt(size_pt)
            r.font.bold = bold
            r.font.color.rgb = RGBColor(*color_rgb)


def parse_int_safe(val, default=0):
    """Konversi nilai bobot atau angka secara aman baik dari int, float, maupun string dengan %."""
    if val is None:
        return default
    try:
        if isinstance(val, (int, float)):
            return int(val)
        cleaned = str(val).replace("%", "").strip()
        return int(cleaned) if cleaned else default
    except Exception:
        return default


def set_cell_background(cell, hex_color="F2F2F2"):
    """Mengatur warna latar belakang cell tabel."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def populate_rps(data: dict, template_path: str, output_path: str) -> str:
    """Mengisi template Format RPS UNIROW dengan data JSON dan menyimpannya ke file baru."""
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template RPS tidak ditemukan: {template_path}")

    doc = docx.Document(template_path)
    identitas = data.get("identitas", {})
    capaian = data.get("capaian_pembelajaran", {})
    cpl_list = capaian.get("cpl", [])
    cpmk_list = capaian.get("cpmk", [])
    sub_cpmk_list = capaian.get("sub_cpmk", [])

    # =========================================================================
    # 1. TABEL 0: Kop, Identitas MK, Otorisasi, Capaian Pembelajaran
    # =========================================================================
    t0 = doc.tables[0]

    # Baris 0: Kop Institusi & Kode Dokumen
    kop_text = f"{identitas.get('universitas', 'UNIVERSITAS PGRI RONGGOLAWE')}\n{identitas.get('fakultas', 'FAKULTAS SAINS DAN TEKNOLOGI')}\n{identitas.get('prodi', 'PROGRAM STUDI S1 INFORMATIKA')}"
    t0.rows[0].cells[1].text = kop_text
    set_cell_font(t0.rows[0].cells[1], name="Calibri", size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    kode_dok_text = f"Kode Dokumen:\n{identitas.get('kode_dokumen', 'RPS-INF-XXX')}"
    t0.rows[0].cells[5].text = kode_dok_text
    set_cell_font(t0.rows[0].cells[5], name="Calibri", size_pt=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Baris 3: Data Identitas Mata Kuliah
    t0.rows[3].cells[0].text = identitas.get("mata_kuliah", "")
    t0.rows[3].cells[1].text = identitas.get("kode_mk", "")
    t0.rows[3].cells[2].text = identitas.get("rumpun_mk", "")
    t0.rows[3].cells[3].text = str(identitas.get("bobot_sks", ""))
    t0.rows[3].cells[4].text = str(identitas.get("semester", ""))
    t0.rows[3].cells[5].text = identitas.get("tgl_penyusunan", "")
    for c_i in range(6):
        set_cell_font(t0.rows[3].cells[c_i], size_pt=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Baris 5: Otorisasi / Pengesahan
    dosen = identitas.get("dosen_pengembang", "")
    koordinator = identitas.get("koordinator_rmk", "")
    kaprodi = identitas.get("ka_prodi", "")
    nidn_kaprodi = identitas.get("nidn_kaprodi", "")

    t0.rows[5].cells[1].text = f"\n\nTTD\n\n{dosen}"
    t0.rows[5].cells[3].text = f"\n\nTTD\n\n{koordinator}"
    t0.rows[5].cells[4].text = f"\n\nTTD\n\n{kaprodi}\nNIDN: {nidn_kaprodi}" if nidn_kaprodi else f"\n\nTTD\n\n{kaprodi}"
    for c_i in [1, 3, 4]:
        set_cell_font(t0.rows[5].cells[c_i], size_pt=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Mengisi Capaian Pembelajaran pada Tabel 0
    # CPL pada Baris 7
    cpl_codes = "\n\n".join([c["kode"] for c in cpl_list])
    cpl_descs = "\n\n".join([f"{c['kode']}: {c['deskripsi']}" for c in cpl_list])
    t0.rows[7].cells[1].text = cpl_codes
    t0.rows[7].cells[2].text = cpl_descs
    set_cell_font(t0.rows[7].cells[1], size_pt=9, bold=True)
    set_cell_font(t0.rows[7].cells[2], size_pt=9)

    # CPMK pada Baris 10
    cpmk_codes = "\n\n".join([f"{c['kode']} ({c.get('kontribusi_persen', 100)}%)" for c in cpmk_list])
    cpmk_descs = "\n\n".join([f"{c['kode']} [induk {c.get('cpl_induk', '')}]: {c['deskripsi']}" for c in cpmk_list])
    t0.rows[10].cells[1].text = cpmk_codes
    t0.rows[10].cells[2].text = cpmk_descs
    set_cell_font(t0.rows[10].cells[1], size_pt=9, bold=True)
    set_cell_font(t0.rows[10].cells[2], size_pt=9)

    # Sub-CPMK pada Baris 13
    sub_codes = "\n\n".join([s["kode"] for s in sub_cpmk_list])
    sub_descs = "\n\n".join([f"{s['kode']}: {s['deskripsi']}" for s in sub_cpmk_list])
    t0.rows[13].cells[1].text = sub_codes
    t0.rows[13].cells[2].text = sub_descs
    set_cell_font(t0.rows[13].cells[1], size_pt=9, bold=True)
    set_cell_font(t0.rows[13].cells[2], size_pt=9)

    # Hapus baris placeholder redundant pada Tabel 0 (baris 14, 11, 8 dari belakang)
    for del_idx in [14, 11, 8]:
        if del_idx < len(t0.rows):
            tr = t0.rows[del_idx]._tr
            tr.getparent().remove(tr)

    # =========================================================================
    # 2. TABEL 1: Korelasi CPL-SubCPMK, Deskripsi, BK, Pustaka, Dosen, Prasyarat
    # =========================================================================
    t1 = doc.tables[1]

    # Mengisi Tabel Bersarang (Nested Table) Korelasi CPL terhadap Sub-CPMK
    nested = t1.rows[1].cells[1].tables[0]

    # Atur Header Nested Table
    # Kolom 0: Sub-CPMK, Kolom 1..4: Kode CPL, Kolom 5: Bobot (%), Kolom 6: Jumlah Minggu
    nested_header = nested.rows[0]
    nested_header.cells[0].text = "Sub-CPMK"
    set_cell_font(nested_header.cells[0], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    for i in range(4):
        col_idx = i + 1
        if i < len(cpl_list):
            nested_header.cells[col_idx].text = cpl_list[i]["kode"]
        else:
            nested_header.cells[col_idx].text = "-"
        set_cell_font(nested_header.cells[col_idx], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    nested_header.cells[5].text = "Bobot (%)"
    nested_header.cells[6].text = "Jumlah Minggu"
    set_cell_font(nested_header.cells[5], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_font(nested_header.cells[6], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Hapus baris data lama pada nested table
    while len(nested.rows) > 1:
        tr = nested.rows[1]._tr
        tr.getparent().remove(tr)

    # Tambahkan baris per Sub-CPMK
    total_bobot = 0
    total_minggu = 0
    for s in sub_cpmk_list:
        nr = nested.add_row()
        nr.cells[0].text = s["kode"]
        set_cell_font(nr.cells[0], size_pt=8.5, bold=True)

        cpmk_induk = s.get("cpmk_induk", "")
        for i in range(4):
            col_idx = i + 1
            if i < len(cpl_list):
                cpl_k = cpl_list[i]["kode"]
                # Cek korelasi apakah CPMK induk terkait dengan CPL ini
                terkait = any(cpl_k in cp.get("cpl_induk", "") for cp in cpmk_list if cp["kode"] in cpmk_induk)
                nr.cells[col_idx].text = "V" if terkait else ""
            else:
                nr.cells[col_idx].text = ""
            set_cell_font(nr.cells[col_idx], size_pt=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

        bobot = parse_int_safe(s.get("bobot_penilaian", 0))
        minggu = parse_int_safe(s.get("jumlah_minggu", 1), default=1)
        total_bobot += bobot
        total_minggu += minggu

        nr.cells[5].text = f"{bobot}%"
        nr.cells[6].text = str(minggu)
        set_cell_font(nr.cells[5], size_pt=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_font(nr.cells[6], size_pt=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Baris Total Nested Table
    nr_tot = nested.add_row()
    nr_tot.cells[0].text = "Total Bobot & Alokasi Minggu"
    set_cell_font(nr_tot.cells[0], size_pt=8.5, bold=True)
    for i in range(1, 5):
        nr_tot.cells[i].text = ""
    nr_tot.cells[5].text = f"{total_bobot}%"
    nr_tot.cells[6].text = f"{total_minggu} Minggu"
    set_cell_font(nr_tot.cells[5], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_font(nr_tot.cells[6], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(nr_tot.cells[0], "E8EEF5")
    set_cell_background(nr_tot.cells[5], "E8EEF5")
    set_cell_background(nr_tot.cells[6], "E8EEF5")

    # Baris 2: Deskripsi Singkat MK
    t1.rows[2].cells[1].text = data.get("deskripsi_singkat_mk", "")
    set_cell_font(t1.rows[2].cells[1], size_pt=9)

    # Baris 3: Bahan Kajian: Materi pembelajaran
    bahan_kajian = data.get("bahan_kajian", [])
    bk_text = "\n".join(bahan_kajian) if isinstance(bahan_kajian, list) else str(bahan_kajian)
    t1.rows[3].cells[1].text = bk_text
    set_cell_font(t1.rows[3].cells[1], size_pt=9)

    # Baris 5: Pustaka Utama
    pustaka = data.get("pustaka", {})
    pustaka_utama = pustaka.get("utama", [])
    pu_text = "\n".join([f"[{i+1}] {p}" for i, p in enumerate(pustaka_utama)])
    t1.rows[5].cells[1].text = pu_text
    set_cell_font(t1.rows[5].cells[1], size_pt=9)

    # Baris 7: Pustaka Pendukung
    pustaka_pendukung = pustaka.get("pendukung", [])
    pp_text = "\n".join([f"[{i+1}] {p}" for i, p in enumerate(pustaka_pendukung)])
    t1.rows[7].cells[1].text = pp_text
    set_cell_font(t1.rows[7].cells[1], size_pt=9)

    # Baris 8: Dosen Pengampu
    t1.rows[8].cells[1].text = identitas.get("dosen_pengampu", "")
    set_cell_font(t1.rows[8].cells[1], size_pt=9)

    # Baris 9: Matakuliah syarat
    t1.rows[9].cells[1].text = identitas.get("matakuliah_syarat", "Tidak Ada")
    set_cell_font(t1.rows[9].cells[1], size_pt=9)

    # =========================================================================
    # 3. TABEL 2: Matriks Pembelajaran 16 Pertemuan
    # =========================================================================
    t2 = doc.tables[2]

    # Hapus baris placeholder lama dari indeks 3 ke atas
    while len(t2.rows) > 3:
        tr = t2.rows[3]._tr
        tr.getparent().remove(tr)

    matriks_16 = data.get("matriks_16_pertemuan", [])
    total_bobot_matriks = 0

    for m in matriks_16:
        row = t2.add_row()
        mg = m.get("minggu", "")
        row.cells[0].text = str(mg)
        row.cells[1].text = m.get("sub_cpmk", "")
        row.cells[2].text = m.get("indikator", "")
        row.cells[3].text = m.get("teknik_kriteria", "")
        row.cells[4].text = m.get("luring", "")
        row.cells[5].text = m.get("daring", "-")
        row.cells[6].text = m.get("materi_pustaka", "")

        bobot = parse_int_safe(m.get("bobot", 0))
        total_bobot_matriks += bobot
        row.cells[7].text = f"{bobot}%"

        # Format cell font
        set_cell_font(row.cells[0], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_font(row.cells[1], size_pt=8.5, bold=True)
        set_cell_font(row.cells[2], size_pt=8.5)
        set_cell_font(row.cells[3], size_pt=8.5)
        set_cell_font(row.cells[4], size_pt=8.5)
        set_cell_font(row.cells[5], size_pt=8.5)
        set_cell_font(row.cells[6], size_pt=8.5)
        set_cell_font(row.cells[7], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Highlight khusus untuk UTS (Mg 8) dan UAS (Mg 16)
        if str(mg) in ["8", "16"] or "UTS" in m.get("sub_cpmk", "") or "UAS" in m.get("sub_cpmk", ""):
            for c in row.cells:
                set_cell_background(c, "FFF2CC")

    # Baris Total Bobot Matriks
    row_total = t2.add_row()
    row_total.cells[0].text = "Total Bobot Penilaian"
    for c_i in range(1, 7):
        row_total.cells[c_i].text = ""
    row_total.cells[7].text = f"{total_bobot_matriks}%"

    set_cell_font(row_total.cells[0], size_pt=9, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_cell_font(row_total.cells[7], size_pt=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c in row_total.cells:
        set_cell_background(c, "D9E1F2")

    # =========================================================================
    # 4. RANCANGAN TUGAS TERSTRUKTUR & RUBRIK PENILAIAN OBE SKALA 1-5
    # =========================================================================
    tugas_list = data.get("tugas_terstruktur", [])
    rubrik_obe = data.get("rubrik_obe", {})

    if tugas_list:
        doc.add_page_break()
        h_tugas = doc.add_heading("RANCANGAN TUGAS TERSTRUKTUR MAHASISWA", level=1)
        h_tugas.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for t in tugas_list:
            p_title = doc.add_paragraph()
            r_title = p_title.add_run(f"TUGAS {t.get('tugas_ke', '')}: {t.get('judul', '').upper()}")
            r_title.bold = True
            r_title.font.size = Pt(11)

            t_table = doc.add_table(rows=6, cols=2)
            t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            t_table.autofit = False

            t_labels = [
                ("Mata Kuliah / SKS", f"{identitas.get('mata_kuliah', '')} / {identitas.get('bobot_sks', '')} SKS"),
                ("Sub-CPMK yang Disasar", t.get("sub_cpmk", "")),
                ("Bentuk Tugas / Alokasi", f"{t.get('jenis', 'Kelompok')} | Waktu: {t.get('alokasi_waktu', '2 Minggu')}"),
                ("Deskripsi Pengerjaan", t.get("deskripsi", "")),
                ("Indikator Penilaian", t.get("indikator", "")),
                ("Bentuk & Format Luaran", f"{t.get('luaran', '')} (Bobot: {t.get('bobot', '')}%)")
            ]

            for r_i, (lbl, val) in enumerate(t_labels):
                t_table.rows[r_i].cells[0].width = Inches(2.2)
                t_table.rows[r_i].cells[1].width = Inches(5.8)
                t_table.rows[r_i].cells[0].text = lbl
                t_table.rows[r_i].cells[1].text = val
                set_cell_font(t_table.rows[r_i].cells[0], size_pt=9.5, bold=True)
                set_cell_font(t_table.rows[r_i].cells[1], size_pt=9.5)
                set_cell_background(t_table.rows[r_i].cells[0], "F2F2F2")

            doc.add_paragraph()  # spacing

    # Tabel Rubrik Penilaian OBE Skala 1-5
    if isinstance(rubrik_obe, dict):
        skala_list = rubrik_obe.get("skala", [])
    elif isinstance(rubrik_obe, list):
        skala_list = rubrik_obe
    else:
        skala_list = []

    if skala_list:
        h_rubrik = doc.add_heading("RUBRIK PENILAIAN OBE SKALA 1–5 (STANDAR APTIKOM & UNIROW)", level=2)
        r_table = doc.add_table(rows=1, cols=4)
        r_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        r_header = r_table.rows[0]
        r_cols = ["Skala", "Kategori Mutu", "Rentang Skor", "Deskripsi Kinerja Ketercapaian"]
        for c_i, name in enumerate(r_cols):
            r_header.cells[c_i].text = name
            set_cell_font(r_header.cells[c_i], size_pt=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_background(r_header.cells[c_i], "D9E1F2")

        for sk in skala_list:
            row = r_table.add_row()
            if isinstance(sk, dict) and "nilai" in sk:
                row.cells[0].text = str(sk.get("nilai", ""))
                row.cells[1].text = sk.get("kategori", "")
                row.cells[2].text = sk.get("rentang", "")
                row.cells[3].text = sk.get("deskripsi", "")
            elif isinstance(sk, dict) and "aspek" in sk:
                row.cells[0].text = "-"
                row.cells[1].text = sk.get("aspek", "")
                row.cells[2].text = "Skala 1-5"
                row.cells[3].text = f"Sangat Baik: {sk.get('skor_5', '')}\nCukup: {sk.get('skor_3', '')}\nKurang: {sk.get('skor_1', '')}"
            else:
                row.cells[0].text = "-"
                row.cells[1].text = str(sk)
                row.cells[2].text = "-"
                row.cells[3].text = "-"

            set_cell_font(row.cells[0], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_font(row.cells[1], size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_font(row.cells[2], size_pt=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_font(row.cells[3], size_pt=8.5)

    # =========================================================================
    # 5. PEMBARUAN LEMBAR VALIDASI AKHIR
    # =========================================================================
    tgl_val = identitas.get("tgl_penyusunan", "September 2026")
    prodi_name = identitas.get("prodi", "S1 Informatika")
    kaprodi_name = identitas.get("ka_prodi", "Nama Kaprodi")
    nidn_kaprodi_str = identitas.get("nidn_kaprodi", "")
    ujm_name = identitas.get("ujm_prodi", "Nama UJM")
    nidn_ujm_str = identitas.get("nidn_ujm", "")

    for p in doc.paragraphs:
        if "RPS ini telah divalidasi pada tanggal" in p.text:
            p.text = f"RPS ini telah divalidasi pada tanggal {tgl_val}"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.bold = True
        elif "(Nama Prodi)" in p.text:
            p.text = f"{prodi_name}\t\t\t\t\t{prodi_name}"
        elif "(Nama Kaprodi)" in p.text:
            p.text = f"{kaprodi_name}\t\t\t\t\t{ujm_name}"
            for r in p.runs:
                r.font.bold = True
                r.font.underline = True
        elif "NIDN/NUPTK" in p.text:
            p.text = f"NIDN: {nidn_kaprodi_str}\t\t\t\t\tNIDN: {nidn_ujm_str}"

    doc.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Generator Dokumen RPS UNIROW Berbasis OBE")
    parser.add_argument("--input", "-i", help="Path file data RPS dalam format JSON")
    parser.add_argument("--output", "-o", help="Path file output .docx")
    parser.add_argument("--template", "-t", help="Path template Format RPS UNIROW.docx")
    parser.add_argument("--example", "-e", action="store_true", help="Gunakan data contoh Rekayasa Perangkat Lunak")
    args = parser.parse_args()

    script_dir = Path(__file__).resolve().parent
    skill_root = script_dir.parent

    # Path template default
    default_template = skill_root / "resources" / "Format_RPS_UNIROW_Template.docx"
    if not default_template.exists():
        default_template = Path("/run/media/andy/512/Uncategories/Skill RPS/Format RPS UNIROW.docx")

    template_path = args.template if args.template else str(default_template)

    if args.example:
        example_json = skill_root / "examples" / "rps_contoh_rekayasa_perangkat_lunak.json"
        output_docx = skill_root / "examples" / "RPS_MK08_Rekayasa_Perangkat_Lunak_UNIROW.docx"
        with open(example_json, "r", encoding="utf-8") as f:
            data = json.load(f)
        out = populate_rps(data, template_path, str(output_docx))
        print(f" Sukses membuat contoh RPS: {out}")
        return

    if not args.input:
        print("Error: Mohon berikan argumen --input <file.json> atau --example.")
        parser.print_help()
        sys.exit(1)

    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    output_path = args.output
    if not output_path:
        mk_kode = data.get("identitas", {}).get("kode_mk", "MK")
        mk_nama = data.get("identitas", {}).get("mata_kuliah", "RPS").replace(" ", "_")
        output_path = f"RPS_{mk_kode}_{mk_nama}_UNIROW.docx"

    out = populate_rps(data, template_path, output_path)
    print(f" Dokumen RPS resmi berhasil digenerasi: {out}")


if __name__ == "__main__":
    main()
